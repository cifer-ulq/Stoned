import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION, WD_ORIENTATION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, color in borders.items():
        if color:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def create_document():
    doc = docx.Document()
    
    # Set page layout to Landscape Letter for large 10-column table
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37) # slate-800
    
    # Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("CARLOS HILADO MEMORIAL STATE UNIVERSITY\nCollege of Computer Studies\n")
    run_inst.font.name = 'Arial'
    run_inst.font.size = Pt(10)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(0x00, 0x59, 0x30) # CHMSU Green
    
    run_sys = title_p.add_run("HIRE ME: A Centralized Platform for Digital Portfolio, Career Job Matching Recommendation, and Mapping\n")
    run_sys.font.name = 'Arial'
    run_sys.font.size = Pt(11)
    run_sys.font.bold = True
    run_sys.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    
    run_doc = title_p.add_run("FUNCTIONAL TEST CASE EVALUATION DOCUMENT")
    run_doc.font.name = 'Arial'
    run_doc.font.size = Pt(13)
    run_doc.font.bold = True
    run_doc.font.color.rgb = RGBColor(0x00, 0x59, 0x30)
    
    # Thin divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_before = Pt(2)
    div_p.paragraph_format.space_after = Pt(6)
    r_div = div_p.add_run("━" * 88)
    r_div.font.size = Pt(8)
    r_div.font.color.rgb = RGBColor(0x00, 0x59, 0x30)
    div_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Evaluator Profile Box (2-column table)
    prof_table = doc.add_table(rows=3, cols=2)
    prof_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prof_table.autofit = False
    
    col_widths = [Inches(5.0), Inches(5.0)]
    for row in prof_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            set_cell_background(cell, "F9FAFB")
            set_cell_borders(cell, top="E5E7EB", bottom="E5E7EB", left="E5E7EB", right="E5E7EB")

    prof_data = [
        ("Evaluator Name (Optional):  ____________________________________", "Role / Occupation:          ____________________________________"),
        ("Years of IT Experience:         ____________________________________", "Organization/Institution: ____________________________________"),
        ("Testing Environment:           Desktop / Laptop (Web Browser)",      "Date of Evaluation:          ____________________________________")
    ]
    for r_idx, (c1, c2) in enumerate(prof_data):
        row = prof_table.rows[r_idx]
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(c1)
        r1.font.size = Pt(9)
        r1.font.bold = ("Evaluator Name" in c1 or "Years of IT" in c1 or "Testing Environment" in c1)
        
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(c2)
        r2.font.size = Pt(9)
        r2.font.bold = ("Role" in c2 or "Organization" in c2 or "Date" in c2)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # General Instructions & Evaluation Scale
    inst_p = doc.add_paragraph()
    inst_p.paragraph_format.space_before = Pt(2)
    inst_p.paragraph_format.space_after = Pt(4)
    run_inst_hdr = inst_p.add_run("General Instructions:\n")
    run_inst_hdr.font.bold = True
    run_inst_hdr.font.size = Pt(10)
    run_inst_hdr.font.color.rgb = RGBColor(0x00, 0x59, 0x30)
    
    inst_body = inst_p.add_run(
        "Please execute the following functional test cases within the \"Hire Me: A Centralized Platform for Digital Portfolio, Career Job Matching Recommendation, and Mapping\" system to rigorously validate its core technical features across all system modules. For each test scenario, verify that all Preconditions are satisfied, strictly execute the Test Steps, and apply the exact Test Data specified. Upon completing each execution, observe the system behavior and record it in the Actual Result column. Finally, evaluate the functional correctness using the binary scale:\n"
        "• Pass: The system executed the scenario completely and correctly, and the actual outcome matches the expected technical result.\n"
        "• Fail: The system failed to execute the scenario, produced runtime/API errors, required unintended manual intervention, or the actual result deviated from the expected outcome."
    )
    inst_body.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Test Cases Execution Checklist (10 Columns)
    headers = [
        "Test Case ID",
        "Test Scenario / Objective",
        "Test Description",
        "Preconditions",
        "Test Steps",
        "Test Data",
        "Expected Result",
        "Actual Result",
        "Status\n(Pass/Fail)",
        "Remarks"
    ]
    
    # 14 Test Cases Data
    test_cases = [
        {
            "id": "TC-01",
            "scenario": "Centralized Recruitment\n(Obj 1.1)",
            "desc": "Verify administrative authentication and secure session initialization for authorized CIER personnel on the centralized web portal.",
            "precond": "CIER Administrator account exists with active role permissions in the database (role = 'admin').",
            "steps": "1. Open web browser and access the administrative login interface (/admin-login.html).\n2. Input administrative email address in the Email field.\n3. Input valid password in the Password field.\n4. Click the \"Sign In\" button (#submit-btn).",
            "data": "Email: admin@chmsu.edu.ph\nPassword: Admin@123",
            "expected": "The system validates credentials via POST /api/auth/admin-login, issues a Sanctum bearer token (admin_token), stores session in localStorage, and redirects the administrator to the CIER Admin Dashboard (/admin/) displaying system statistics and management modules."
        },
        {
            "id": "TC-02",
            "scenario": "Centralized Recruitment\n(Obj 1.1)",
            "desc": "Verify access denial, credential rejection, and error prompt display when attempting authentication with an incorrect password.",
            "precond": "A registered partner company account exists in the database with role 'company'.",
            "steps": "1. Open the main portal login page (/ or /login/).\n2. Ensure the \"Sign In\" tab (#tab-signin) is selected.\n3. Enter registered company email address.\n4. Enter an invalid/incorrect password.\n5. Click the \"Sign In\" button.",
            "data": "Email: company@gmail.com\nPassword: WrongPass!",
            "expected": "The backend API (POST /api/auth/login) rejects authentication with HTTP 401 Unauthorized, blocks token issuance, and displays the inline error message: \"Invalid credentials.\" next to the input field while remaining on the sign-in form."
        },
        {
            "id": "TC-03",
            "scenario": "Digital Portfolio\n(Obj 1.3)",
            "desc": "Verify dynamic saving of academic background, technical competencies, and automated real-time generation of the standardized 1-page ATS-ready digital resume.",
            "precond": "Student is authenticated and logged into the student portal with an active session.",
            "steps": "1. Navigate to \"My Portfolio\" (#portfolio) from sidebar.\n2. Under \"About\" tab, click \"+ Add Education\", input institution, degree, year level, and GPA, then click \"Add Education\".\n3. Under Skills section, click \"+ Add Skill\", select category, enter technical skill, select proficiency level (0–100 scale), and click \"Save Skill\".\n4. Under \"Achievements\" tab, upload certificate file.\n5. Switch to the \"Resume\" tab to inspect generated resume.",
            "data": "User: student@gmail.com\nSchool: Carlos Hilado Memorial State University\nDegree: BS in Information Technology\nSkill: PHP / Laravel (Proficiency: 85%)\nFile: certificate.pdf",
            "expected": "System persists credentials in database (student_educations, student_skills, student_achievements), dynamically compiles them into the standardized executive 1-Page ATS-ready Resume view on the canvas, renders verified headers, and activates \"Fullscreen / Print\" and \"Download PDF\" buttons."
        },
        {
            "id": "TC-04",
            "scenario": "Digital Portfolio\n(Obj 1.3)",
            "desc": "Verify form input validation and error rejection when attempting to save portfolio education records with missing mandatory fields.",
            "precond": "Student is authenticated and logged into the student portal.",
            "steps": "1. Navigate to \"My Portfolio\" (#portfolio) and stay on the \"About\" tab.\n2. Click \"+ Add Education\" to open modal dialog.\n3. Leave the mandatory fields (\"School / University *\" and \"Degree / Program *\") completely blank.\n4. Click the \"Add Education\" / \"Save Changes\" button (.pfm-save-btn).",
            "data": "School / University: [Blank]\nDegree / Program: [Blank]",
            "expected": "System prevents form submission, halts backend API dispatch, enforces required field validation, and displays the inline error message: \"School and degree are required.\" inside the modal dialog while preserving the modal open for user correction."
        },
        {
            "id": "TC-05",
            "scenario": "External Job Retrieval\n(Obj 1.4)",
            "desc": "Verify real-time external job aggregation and display from third-party RESTful APIs based on candidate technical profile competencies.",
            "precond": "User is logged into portal, has registered skills in profile, and active external internet connection is available.",
            "steps": "1. Navigate to \"Career Opportunities / Jobs\" (#jobs) from navigation menu.\n2. In the job filter tab bar, click the \"External Jobs\" tab (globe icon).\n3. Wait for the asynchronous RESTful API fetch to resolve.",
            "data": "Candidate Skills: Programmer / Developer / PHP / React\nEndpoint: GET /api/student/external-jobs (RapidAPI JSearch RESTful API)",
            "expected": "System initiates an HTTP request to the third-party RESTful API (RapidAPI JSearch), parses live employment listings, and dynamically populates cards displaying job title, employer, location, salary tier, \"External\" pill badge, and the \"Apply on Site\" external application modal."
        },
        {
            "id": "TC-06",
            "scenario": "Skill Matching\n(Obj 1.2)",
            "desc": "Verify automated candidate skill matching algorithm accurately computes percentage compatibility score and assigns recommendation tier based on profile competencies versus vacancy requirements.",
            "precond": "Graduate/Student has populated technical skills in profile; company has an active job posting with defined required skills.",
            "steps": "1. Log into portal as graduate or student with registered technical competencies (e.g. PHP, MySQL, JavaScript).\n2. Navigate to \"Career Opportunities\" (#jobs).\n3. View active job vacancies in the primary opportunity feed.\n4. Inspect the match score badge, match explanation insight, and recommendation tier for a target vacancy.",
            "data": "Candidate Skills: PHP, MySQL, JavaScript\nJob Listing Required Skills: PHP, MySQL, JavaScript, HTML (3 out of 4 matching = 75% compatibility)",
            "expected": "The system's weighted matching algorithm calculates compatibility (score = 75%), displays the badge \"75% Match\", categorizes position as \"Highly Recommended\" (score ≥ 71%) with the award badge, highlights matching skill tags in green, and prioritizes listing in the feed."
        },
        {
            "id": "TC-07",
            "scenario": "Skill Matching\n(Obj 1.2)",
            "desc": "Verify algorithm classification, warning banner generation, and feed filtering when a candidate possesses zero matching skills for an employer vacancy.",
            "precond": "Candidate profile skills have zero overlap with the required skills of an active job vacancy.",
            "steps": "1. Log into portal and navigate to \"Career Opportunities\" (#jobs).\n2. Locate an active job vacancy where none of the required skills match candidate competencies.\n3. Observe displayed match badge and recommendation tier.\n4. Click \"Apply Now\" on the card to inspect modal behavior.\n5. Click the \"Top Matches\" filter tab.",
            "data": "Candidate Skills: Graphic Design, Photoshop\nJob Required Skills: Python, Django, PostgreSQL\n(0 matching skills = 0% fit)",
            "expected": "The matching engine computes a \"0% Match\", assigns the \"Not Recommended\" classification badge (score ≤ 25%), displays an advisory insight (\"Host companies prioritize applicants with matching skills\"), renders a cautionary warning banner inside the application modal, and completely filters out the job from the \"Top Matches\" feed."
        },
        {
            "id": "TC-08",
            "scenario": "Graduate Screening\n(Obj 1.5)",
            "desc": "Verify candidate job application submission workflow, automatic digital portfolio attachment, and application pipeline status tracking.",
            "precond": "Graduate/Student is authenticated with a complete digital profile and viewing an active internal job vacancy.",
            "steps": "1. Navigate to \"Career Opportunities\" (#jobs).\n2. Locate target vacancy and click the \"Apply Now\" button (.jm-card__apply).\n3. In the Application Modal, review automated profile attachment notice.\n4. Enter optional text in \"Cover Letter / Introduction\".\n5. Click the \"Submit Application\" button.",
            "data": "Job Vacancy: Active company job listing ID\nCover Letter: \"I am interested in applying for this position.\"",
            "expected": "The system executes POST /api/jobseeker/apply/{jobId}, persists record in job_applications table with status \"applied\", displays a confirmation toast, changes the job card button to a checked \"Applied\" state, and updates the application list on the candidate's \"Applications\" page."
        },
        {
            "id": "TC-09",
            "scenario": "Interview Scheduling\n(Obj 1.5)",
            "desc": "Verify partner company can configure and dispatch an interview schedule appointment for an applicant, updating the recruitment pipeline status.",
            "precond": "Partner company HR is authenticated, and candidate application exists in the applicant review pipeline.",
            "steps": "1. Log into Company Portal (/company/) and navigate to \"Applicants & Candidates\" (#applicants).\n2. Locate candidate card and click \"Schedule Interview\" (.ap-btn-status).\n3. In Schedule Modal, select Interview Format (\"Online Video Meeting\"), Interview Type (\"Technical Interview\"), Date, Time, Duration (\"45 min\"), and enter Meeting Link.\n4. Click \"Schedule Interview\" (#sched-submit).",
            "data": "Format: Online Video Meeting\nType: Technical Interview\nDate: Tomorrow's Date\nTime: 10:00 AM\nMeeting Link: https://meet.google.com/abc-defg-hij",
            "expected": "The system calls POST /api/company/applications/{id}/interview, stores appointment details in interviews table, updates pipeline status to \"interview\" / \"interview_scheduled\", dispatches an interview notification to the candidate, and displays the event on the interview calendar."
        },
        {
            "id": "TC-10",
            "scenario": "OJT Mapping\n(Obj 1.6)",
            "desc": "Verify OJT supervisor monitoring dashboard renders interactive geospatial Leaflet map plotting verified host training sites and active student deployment pins.",
            "precond": "OJT Supervisor/Coordinator is authenticated and student trainees are actively deployed at partner company locations with geographic coordinates.",
            "steps": "1. Log into the Supervisor Portal (/supervisors/).\n2. Navigate to \"OJT Monitoring & Site Geofencing\" (#/monitoring) from sidebar.\n3. Locate a host company site card with active trainees.\n4. Click \"View Site Map & Geofence\" to launch the interactive GIS map.\n5. Click on host site marker and student deployment markers.",
            "data": "Supervisor: supervisor@gmail.com\nHost Training Site: Verified company coordinates (Lat: 10.6975, Lng: 122.9570)",
            "expected": "System loads monitoring data via GET /api/supervisor/monitoring, initializes the Leaflet map centered at host site coordinates, renders the host company building pin with name and address popup, plots student location pins corresponding to verified attendance check-ins, and displays the 100-meter circular geofence boundary."
        },
        {
            "id": "TC-11",
            "scenario": "OJT Geofencing\n(Obj 1.6)",
            "desc": "Verify automated geofence radius verification accepts student attendance punch when device GPS is within the 100-meter perimeter of host company site.",
            "precond": "Trainee is logged into portal, has an active OJT deployment, and device GPS location is within ≤ 100 meters of host training site.",
            "steps": "1. Access Student Portal (/main/) and navigate to \"OJT Time Tracker\" (#ojt-tracker).\n2. Confirm device GPS location permissions are enabled.\n3. Click the active \"Time In\" button for the morning or afternoon shift session.",
            "data": "Device GPS: Latitude & Longitude within 100m radius of host company\nShift: Morning session (e.g. 08:00 AM)",
            "expected": "System calculates Haversine distance (≤ 100m), validates location as \"In Site\", saves attendance record via POST /api/student/ojt-tracker/log-in with status \"valid\", displays success confirmation, renders a green \"In Site\" badge with verified distance on the tracker dashboard, and increments daily rendered hours."
        },
        {
            "id": "TC-12",
            "scenario": "OJT Geofencing\n(Obj 1.6)",
            "desc": "Verify system detects out-of-bounds geolocation, rejects valid attendance credit, and displays a warning dialog when device GPS is outside the 100-meter geofence.",
            "precond": "Trainee is logged into student portal with an active deployment, but device GPS is situated beyond the 100-meter radius (> 100m) of host company site.",
            "steps": "1. Navigate to \"OJT Time Tracker\" (#ojt-tracker).\n2. Attempt to click the \"Time In\" button while physically outside the company geofence boundary.\n3. Observe system validation response and modal alerts.",
            "data": "Device GPS: Coordinates situated 350 meters away from host company site boundary",
            "expected": "System computes Haversine distance (> 100m), flags location validity as \"Too Far\" (status: not_valid), triggers alert modal: \"Location Are Too Far - You are [X] m away from the OJT site. Your log-in was recorded but marked as Too Far\", renders a red/amber \"Too Far\" badge, and withholds valid attendance hours credit pending supervisor audit."
        },
        {
            "id": "TC-13",
            "scenario": "Centralized Recruitment\n(Obj 1.1)",
            "desc": "Verify CIER Administrator can upload notarized bilateral Memorandum of Agreement (MOA) document, set validity duration, and activate partner company status.",
            "precond": "CIER Administrator is logged into Admin Portal, and a partner company exists with \"Pending\" or \"MOA Requested\" status.",
            "steps": "1. Log into Admin Portal (/admin/) and navigate to \"Company Partners\" (#companies).\n2. Click on a partner company entry to open the Company Details Modal.\n3. In \"Upload / Update MOA Form\", input MOA Start Date and Expiration Date.\n4. Attach a valid notarized bilateral MOA document file (.pdf or .docx).\n5. Click the \"Approve & Activate MOA\" / \"Send & Activate MOA\" button (#btn-submit-moa).",
            "data": "Admin: admin@chmsu.edu.ph\nMOA Dates: Start Date: Current Date; Expiration Date: +2 Years\nFile: CHMSU_Company_Bilateral_MOA.pdf",
            "expected": "System executes POST /api/admin/companies/{id}/moa, uploads file to moa-documents storage, updates company and MOA status to \"Active\", dispatches an automated approval notification to the company account, and immediately enables company permissions to post OJT slots and recruitment listings."
        },
        {
            "id": "TC-14",
            "scenario": "Graduate Screening\n(Obj 1.5)",
            "desc": "Verify partner company can initiate real-time direct chat messaging with an applicant and automatically advance candidate review status in the screening pipeline.",
            "precond": "Partner company HR is logged in, and applicant record exists in the candidate screening pipeline.",
            "steps": "1. In Company Portal, navigate to \"Applicants & Candidates\" (#applicants).\n2. Locate target applicant card and click \"View Profile\" (transitioning status from \"applied\" to \"reviewed\").\n3. Click the \"Chat\" button (.ap-btn-chat) on the applicant card.\n4. In the Chat widget/drawer, input direct inquiry message text into the message box.\n5. Click the Send button or press Enter.",
            "data": "Message: \"Hello, we are reviewing your application for the Developer role.\"\nConversation Key: job_app_{id}",
            "expected": "System transmits message via POST /api/chat/{key}/messages, records message in chat_messages table, updates applicant status to \"reviewed\", delivers an unread chat notification to the applicant's account, and renders the message bubble in real time within the active conversation window."
        }
    ]

    # Create table: 1 header row + 14 data rows = 15 rows
    table = doc.add_table(rows=len(test_cases) + 1, cols=10)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Define exact widths for landscape (total 10.0 inches)
    widths = [
        Inches(0.65), # TC ID
        Inches(1.05), # Scenario / Objective
        Inches(1.40), # Description
        Inches(0.95), # Preconditions
        Inches(1.45), # Steps
        Inches(1.05), # Test Data
        Inches(1.55), # Expected Result
        Inches(0.65), # Actual Result
        Inches(0.60), # Status
        Inches(0.65)  # Remarks
    ]

    # Style Header Row
    hdr_row = table.rows[0]
    # Set header repeat across pages
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

    for idx, name in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = widths[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, "005930") # CHMSU Green
        set_cell_margins(cell, top=100, bottom=100, left=60, right=60)
        set_cell_borders(cell, top="003E22", bottom="003E22", left="004D2A", right="004D2A")
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(name)
        run.font.name = 'Arial'
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Populate Data Rows
    for r_idx, tc in enumerate(test_cases):
        row = table.rows[r_idx + 1]
        
        # CantSplit row across page break for clean presentation
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))

        bg_color = "F9FAFB" if (r_idx % 2 == 1) else "FFFFFF"

        row_data = [
            tc["id"],
            tc["scenario"],
            tc["desc"],
            tc["precond"],
            tc["steps"],
            tc["data"],
            tc["expected"],
            "", # Actual Result blank for evaluator input
            "", # Status blank
            ""  # Remarks blank
        ]

        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=50, right=50)
            set_cell_borders(cell, top="D1D5DB", bottom="D1D5DB", left="E5E7EB", right="E5E7EB")

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05

            # Alignment rules
            if c_idx in [0, 8]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(8.0)
            if c_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x00, 0x59, 0x30)
            elif c_idx == 1:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Quantitative Summary & Rating Interpretation Table
    doc.add_paragraph().paragraph_format.space_before = Pt(2)
    sum_title = doc.add_paragraph()
    sum_title.paragraph_format.space_after = Pt(2)
    r_st = sum_title.add_run("Test Case Evaluation Summary & Interpretation Scale")
    r_st.font.bold = True
    r_st.font.size = Pt(10)
    r_st.font.color.rgb = RGBColor(0x00, 0x59, 0x30)

    sum_table = doc.add_table(rows=6, cols=3)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_table.autofit = False
    s_widths = [Inches(2.5), Inches(2.0), Inches(5.5)]

    sum_headers = ["Percentage of Passed Test Cases", "Descriptive Rating", "Technical Interpretation"]
    s_hdr_row = sum_table.rows[0]
    for i, h in enumerate(sum_headers):
        c = s_hdr_row.cells[i]
        c.width = s_widths[i]
        set_cell_background(c, "005930")
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        set_cell_borders(c, top="003E22", bottom="003E22", left="004D2A", right="004D2A")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    scale_rows = [
        ("96% – 100%", "Excellent", "The system passed all test cases with optimal technical precision and zero runtime errors."),
        ("90% – 95%",  "Very Good",  "The system passed most test cases with only minor cosmetic or non-critical interface issues."),
        ("85% – 89%",  "Good",       "The system functions properly but requires minor adjustments in specific workflow modules."),
        ("80% – 84%",  "Fair",       "The system meets baseline functional requirements but several operational defects require correction."),
        ("Below 80%",  "Poor",       "The system failed significant functional test cases and requires critical debugging and revision.")
    ]

    for idx, (p_range, desc, interp) in enumerate(scale_rows):
        row = sum_table.rows[idx + 1]
        bg = "F9FAFB" if idx % 2 == 1 else "FFFFFF"
        for c_idx, text in enumerate([p_range, desc, interp]):
            c = row.cells[c_idx]
            c.width = s_widths[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, top=40, bottom=40, left=80, right=80)
            set_cell_borders(c, top="E5E7EB", bottom="E5E7EB", left="E5E7EB", right="E5E7EB")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if c_idx == 1:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Comments & Suggestions section
    comm_title = doc.add_paragraph()
    comm_title.paragraph_format.space_before = Pt(4)
    comm_title.paragraph_format.space_after = Pt(2)
    r_ct = comm_title.add_run("Evaluator Technical Comments & System Suggestions:")
    r_ct.font.bold = True
    r_ct.font.size = Pt(10)
    r_ct.font.color.rgb = RGBColor(0x00, 0x59, 0x30)

    # Comments Box Table (1 row, 1 col)
    comm_box = doc.add_table(rows=1, cols=1)
    comm_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = comm_box.rows[0].cells[0]
    c_cell.width = Inches(10.0)
    set_cell_background(c_cell, "FAFAFA")
    set_cell_margins(c_cell, top=120, bottom=120, left=120, right=120)
    set_cell_borders(c_cell, top="D1D5DB", bottom="D1D5DB", left="D1D5DB", right="D1D5DB")
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(50) # blank space for hand-writing
    r_cph = cp.add_run("[ Please write qualitative feedback, detected software bugs, algorithmic observations, or mapping remarks below ]\n\n\n\n\n")
    r_cph.font.size = Pt(8.5)
    r_cph.font.italic = True
    r_cph.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Signatures block (2 columns: Evaluator Signature and Capstone Adviser Signature)
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    sig_widths = [Inches(5.0), Inches(5.0)]

    for i in range(2):
        c = sig_table.rows[0].cells[i]
        c.width = sig_widths[i]
        set_cell_borders(c, top=None, bottom=None, left=None, right=None)

    # Left cell: Evaluator
    p_left = sig_table.rows[0].cells[0].paragraphs[0]
    p_left.paragraph_format.space_before = Pt(10)
    p_left.paragraph_format.space_after = Pt(2)
    r_l1 = p_left.add_run("Evaluated by:\n\n\n_________________________________________________\n")
    r_l1.font.size = Pt(9.5)
    r_l2 = p_left.add_run("Signature over Printed Name of IT Professional Evaluator\nDate: ________________________")
    r_l2.font.size = Pt(9)
    r_l2.font.bold = True

    # Right cell: Capstone Adviser
    p_right = sig_table.rows[0].cells[1].paragraphs[0]
    p_right.paragraph_format.space_before = Pt(10)
    p_right.paragraph_format.space_after = Pt(2)
    r_r1 = p_right.add_run("Reviewed and Endorsed by:\n\n\n")
    r_r1.font.size = Pt(9.5)
    r_r2 = p_right.add_run("JAYRELLE B. SY, Ph.D.\n")
    r_r2.font.size = Pt(10)
    r_r2.font.bold = True
    r_r2.font.color.rgb = RGBColor(0x00, 0x59, 0x30)
    r_r3 = p_right.add_run("Capstone Adviser\nDate: ________________________")
    r_r3.font.size = Pt(9)

    out_path = r"C:\Users\Asus'\Desktop\crapstone\Functional_Test_Case_Evaluation_Document_HireMe.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_document()
